"""The project report: its content comes from the engines, and every format renders it."""
import io
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from causal_engine.diagnosis import ResearchDesignInput, diagnose
from causal_engine.recommendation import recommend
from estimators.did import estimate_did
from estimators.matching import estimate_matching
from report_engine.builder import Code, Heading, Paragraph, Table, build_report
from report_engine.render import latex_escape, pdf_safe, render_report, to_latex, to_markdown
from robustness_engine.stress_test import run_stress_test
from simulation_engine.dgp import VirtualWorldConfig, generate
from simulation_engine.method_worlds import MatchingWorldConfig, matching_world
from utils.real_data import load_card_krueger

APP = str(Path(__file__).resolve().parent.parent / "app.py")


def _virtual_state():
    design = ResearchDesignInput(research_question="Do mines raise child mortality? 50% & $_#", treatment_timing="staggered")
    diagnosis = diagnose(design)
    cfg = VirtualWorldConfig(seed=4, n_units=60, staggered_adoption=True)
    df, truth = generate(cfg)
    m_df, m_truth = matching_world(MatchingWorldConfig(seed=1, n=400))
    return {"design": design, "diagnosis": diagnosis, "recommendations": recommend(design, diagnosis),
            "virtual_df": df, "virtual_truth": truth, "virtual_cfg": cfg,
            "method_data": {"matching": {"df": m_df, "truth": m_truth, "source": "virtual",
                                         "result": estimate_matching(m_df, covariate_cols=m_truth["roles"]["covariates"])},
                            "rdd": {"df": None, "result": None}}}


STATE = _virtual_state()
REPORT_EN = build_report(STATE, "en")


def _headings(report):
    return [b.text for b in report.blocks if isinstance(b, Heading)]


def _value(report, label):
    """The value next to `label` in the first two-column table that has it."""
    for block in report.blocks:
        if isinstance(block, Table):
            for row in block.rows:
                if row[0] == label:
                    return row[1]
    raise KeyError(label)


def test_report_has_every_section_and_the_engines_numbers():
    heads = _headings(REPORT_EN)
    for expected in ("Research question and design", "Research Design Diagnosis", "Data", "Estimation",
                     "Difference-in-Differences", "Break My Design", "Robustness Engine", "Reproducible code"):
        assert any(expected in h for h in heads), (expected, heads)
    did = estimate_did(STATE["virtual_df"])
    assert _value(REPORT_EN, "Estimated ATT") == f"{did.att:.4f}"
    assert _value(REPORT_EN, "True ATT") == f"{STATE['virtual_truth']['true_att']:.4f}"
    stress = run_stress_test(STATE["virtual_df"], treatment_period=STATE["virtual_cfg"].treatment_period,
                             ids_encode_neighbourhood=True)
    strength = next(b.text for b in REPORT_EN.blocks if isinstance(b, Paragraph) and b.text.startswith("Identification"))
    assert f"{stress.identification_strength:.0f} / 100" in strength
    assert any("Callaway" in h or "Staggered" in h or "robust" in h.lower() for h in heads)   # staggered world
    matching = STATE["method_data"]["matching"]["result"]
    assert f"{matching.estimate:.4f}" in [row[1] for b in REPORT_EN.blocks if isinstance(b, Table) for row in b.rows]
    assert [b.language for b in REPORT_EN.blocks if isinstance(b, Code)] == ["python", "r", "stata"]


def test_report_follows_the_interface_language():
    fr = build_report(STATE, "fr")
    assert fr.title == "Rapport Causal Lab" and "Données" in _headings(fr)
    assert any(isinstance(b, Paragraph) and b.text.startswith("Méthode principale : ") for b in fr.blocks)
    assert len(fr.blocks) == len(REPORT_EN.blocks)


def test_report_without_data_or_question_says_so():
    report = build_report({}, "en")
    texts = [b.text for b in report.blocks if isinstance(b, Paragraph)]
    assert "No research question was entered." in texts
    assert "Estimation" not in _headings(report) and "Reproducible code" in _headings(report)


def test_real_data_report_has_no_true_effect():
    report = build_report({"uploaded_df": load_card_krueger(), "uploaded_file_id": "card_krueger_1994"}, "en")
    assert _value(report, "Estimated ATT") == "2.7500"
    try:
        _value(report, "True ATT")
    except KeyError:
        pass
    else:
        raise AssertionError("real data have no true effect")
    assert any("Card & Krueger" in b.text for b in report.blocks if isinstance(b, Paragraph))


def test_markdown_escapes_table_pipes():
    md = to_markdown(REPORT_EN)
    assert md.startswith("# Causal Lab report") and "```python" in md
    table = Table(["a|b"], [["x|y"]])
    REPORT_EN.blocks.append(table)
    try:
        assert "| a\\|b |" in to_markdown(REPORT_EN) and "| x\\|y |" in to_markdown(REPORT_EN)
    finally:
        REPORT_EN.blocks.remove(table)


def test_latex_escapes_special_characters():
    assert latex_escape("50% & $_#{}~^\\") == r"50\% \& \$\_\#\{\}\textasciitilde{}\textasciicircum{}\textbackslash{}"
    assert latex_escape("Δ ≤ ± — →") == r"$\Delta$ $\leq$ $\pm$ --- $\rightarrow$"
    tex = to_latex(build_report(STATE, "fr"))
    assert tex.startswith(r"\documentclass") and tex.rstrip().endswith(r"\end{document}")
    assert r"\usepackage[french]{babel}" in tex and r"50\% \& \$\_\#" in tex


def test_latex_compiles_when_a_tex_engine_is_installed():
    # Checked on Windows with MiKTeX; skipped where no pdflatex is installed (CI).
    if shutil.which("pdflatex") is None:
        print("pdflatex not installed: LaTeX compilation not checked")
        return
    with tempfile.TemporaryDirectory() as folder:
        for lang in ("en", "fr"):
            path = Path(folder) / f"report_{lang}.tex"
            path.write_bytes(render_report(build_report(STATE, lang), "latex"))
            run = subprocess.run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", path.name],
                                 cwd=folder, capture_output=True, text=True, timeout=600)
            assert run.returncode == 0, [line for line in run.stdout.splitlines() if line.startswith("!")]
            assert (Path(folder) / f"report_{lang}.pdf").stat().st_size > 10_000


def test_word_document_holds_every_heading_and_table():
    from docx import Document
    doc = Document(io.BytesIO(render_report(REPORT_EN, "docx")))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings == _headings(REPORT_EN)
    assert len(doc.tables) == sum(isinstance(b, Table) for b in REPORT_EN.blocks)
    assert doc.tables[0].rows[0].cells[0].text == "Item"


def test_pdf_is_a_pdf_and_keeps_only_drawable_characters():
    for lang in ("en", "fr"):
        data = render_report(build_report(STATE, lang), "pdf")
        assert data.startswith(b"%PDF-") and data.rstrip().endswith(b"%%EOF") and len(data) > 5000
    assert pdf_safe("Δ ≤ 5 − 1 → été «ok» — 漢") == "Delta  <= 5 - 1 -> été «ok» — ?"


def test_report_page_generates_a_download():
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(APP, default_timeout=300)
    for key, value in STATE.items():
        at.session_state[key] = value
    at.session_state["page"] = "report"
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    assert not at.get("download_button")
    at.button[0].click().run()
    assert not at.exception, [e.value for e in at.exception]
    assert at.session_state["report_output"][1].startswith(b"%PDF-")
    assert at.get("download_button")


if __name__ == "__main__":
    for name, test in list(globals().items()):
        if name.startswith("test_"):
            test()
    print("test_report: OK")
