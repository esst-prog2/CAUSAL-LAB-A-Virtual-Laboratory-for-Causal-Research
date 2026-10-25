"""
Render a `report_engine.builder.Report` as Markdown, LaTeX, Word or PDF.

Markdown and LaTeX are text. Word uses python-docx and PDF uses reportlab,
both pure Python. The PDF uses the standard Helvetica and Courier fonts, which
cover the Windows-1252 character set; the few other characters in the
interface texts (Δ, −, →, ≤, ≥) are spelled out, and any remaining one
(possible in a research question) is replaced by "?".
"""
from __future__ import annotations

import io
import re

from report_engine.builder import Code, Heading, Paragraph, Report, Table

FORMATS = {               # format id: (file extension, MIME type)
    "pdf": (".pdf", "application/pdf"),
    "docx": (".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "latex": (".tex", "application/x-tex"),
    "markdown": (".md", "text/markdown"),
}
_SPELLED_OUT = {"Δ": "Delta ", "−": "-", "→": "->", "≤": "<=", "≥": ">="}


def render_report(report: Report, fmt: str) -> bytes:
    if fmt == "markdown":
        return to_markdown(report).encode("utf-8")
    if fmt == "latex":
        return to_latex(report).encode("utf-8")
    if fmt == "docx":
        return to_docx(report)
    if fmt == "pdf":
        return to_pdf(report)
    raise KeyError(fmt)


# --------------------------------------------------------------------------
# Markdown
# --------------------------------------------------------------------------
def _md_cell(text: str) -> str:
    return str(text).replace("|", "\\|").replace("\n", " ")


def to_markdown(report: Report) -> str:
    out = [f"# {report.title}", "", f"*{report.subtitle}*", ""]
    for block in report.blocks:
        if isinstance(block, Heading):
            out += ["#" * (block.level + 1) + " " + block.text, ""]
        elif isinstance(block, Paragraph):
            out += [f"*{block.text}*" if block.emphasis else block.text, ""]
        elif isinstance(block, Table):
            out.append("| " + " | ".join(_md_cell(c) for c in block.columns) + " |")
            out.append("|" + "|".join("---" for _ in block.columns) + "|")
            out += ["| " + " | ".join(_md_cell(c) for c in row) + " |" for row in block.rows]
            out.append("")
        elif isinstance(block, Code):
            out += [f"```{block.language}", block.text.rstrip("\n"), "```", ""]
    return "\n".join(out)


# --------------------------------------------------------------------------
# LaTeX
# --------------------------------------------------------------------------
_LATEX_SPECIAL = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
                  "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}",
                  "Δ": r"$\Delta$", "−": r"$-$", "→": r"$\rightarrow$", "≤": r"$\leq$", "≥": r"$\geq$",
                  "—": "---", "–": "--", "±": r"$\pm$", "«": r"\guillemotleft{}", "»": r"\guillemotright{}"}
_LATEX_RE = re.compile("|".join(re.escape(k) for k in _LATEX_SPECIAL))


def latex_escape(text: str) -> str:
    return _LATEX_RE.sub(lambda m: _LATEX_SPECIAL[m.group(0)], str(text))


def to_latex(report: Report) -> str:
    babel = "french" if report.lang == "fr" else "english"
    out = [r"\documentclass[11pt,a4paper]{article}", r"\usepackage[utf8]{inputenc}", r"\usepackage[T1]{fontenc}", r"\usepackage{textcomp}",
           rf"\usepackage[{babel}]{{babel}}", r"\usepackage{lmodern}", r"\usepackage[margin=2.2cm]{geometry}",
           r"\usepackage{longtable,booktabs,array}", r"\usepackage{fancyvrb}",
           r"\setlength{\parindent}{0pt}\setlength{\parskip}{0.5em}",
           rf"\title{{{latex_escape(report.title)}}}", rf"\date{{{latex_escape(report.subtitle)}}}", r"\author{}",
           r"\begin{document}", r"\maketitle"]
    for block in report.blocks:
        if isinstance(block, Heading):
            out.append((r"\section*{" if block.level == 1 else r"\subsection*{") + latex_escape(block.text) + "}")
        elif isinstance(block, Paragraph):
            text = latex_escape(block.text)
            out += [r"\emph{" + text + "}" if block.emphasis else text, ""]
        elif isinstance(block, Table):
            n = len(block.columns)
            width = f"{0.9 / n:.3f}"
            out.append(r"\begin{longtable}{" + "".join(rf"p{{{width}\linewidth}}" for _ in range(n)) + "}")
            out.append(r"\toprule")
            out.append(" & ".join(r"\textbf{" + latex_escape(c) + "}" for c in block.columns) + r" \\")
            out.append(r"\midrule\endhead")
            out += [" & ".join(latex_escape(c) for c in row) + r" \\" for row in block.rows]
            out += [r"\bottomrule", r"\end{longtable}"]
        elif isinstance(block, Code):
            # Verbatim keeps every character; only the few outside T1 are spelled out.
            code = block.text.rstrip("\n")
            for char, spelled in _SPELLED_OUT.items():
                code = code.replace(char, spelled)
            out += [r"\begin{Verbatim}[fontsize=\footnotesize,frame=single]", code, r"\end{Verbatim}"]
    out.append(r"\end{document}")
    return "\n".join(out) + "\n"


# --------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------
def to_docx(report: Report) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(report.title, level=0)
    doc.add_paragraph(report.subtitle).runs[0].italic = True
    for block in report.blocks:
        if isinstance(block, Heading):
            doc.add_heading(block.text, level=block.level)
        elif isinstance(block, Paragraph):
            run = doc.add_paragraph().add_run(block.text)
            run.italic = block.emphasis
        elif isinstance(block, Table):
            table = doc.add_table(rows=1, cols=len(block.columns))
            table.style = "Light Grid Accent 1"
            for cell, text in zip(table.rows[0].cells, block.columns):
                cell.text = str(text)
            for row in block.rows:
                for cell, text in zip(table.add_row().cells, row):
                    cell.text = str(text)
            doc.add_paragraph()
        elif isinstance(block, Code):
            for line in block.text.rstrip("\n").split("\n"):
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(0)
                run = paragraph.add_run(line)
                run.font.name, run.font.size = "Courier New", Pt(8)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def pdf_safe(text: str) -> str:
    """Text the standard PDF fonts can draw (Windows-1252)."""
    for char, spelled in _SPELLED_OUT.items():
        text = text.replace(char, spelled)
    return text.encode("cp1252", errors="replace").decode("cp1252")


def to_pdf(report: Report) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph as P, Preformatted, SimpleDocTemplate, Spacer, Table as T, TableStyle

    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=8, leading=10)
    head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold")
    code_style = ParagraphStyle("code", parent=styles["Code"], fontSize=7, leading=8.5)
    text = lambda s, style: P(escape(pdf_safe(str(s))), style)              # noqa: E731

    story = [text(report.title, styles["Title"]), text(report.subtitle, styles["Italic"]), Spacer(1, 0.4 * cm)]
    width = A4[0] - 4 * cm
    for block in report.blocks:
        if isinstance(block, Heading):
            story.append(text(block.text, styles["Heading1" if block.level == 1 else "Heading2"]))
        elif isinstance(block, Paragraph):
            story.append(text(block.text, styles["Italic" if block.emphasis else "BodyText"]))
        elif isinstance(block, Table):
            data = [[text(c, head) for c in block.columns]] + [[text(c, cell) for c in row] for row in block.rows]
            table = T(data, colWidths=[width / len(block.columns)] * len(block.columns), repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE7F0")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA9B8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story += [table, Spacer(1, 0.3 * cm)]
        elif isinstance(block, Code):
            story.append(Preformatted(pdf_safe(block.text.rstrip("\n")), code_style, maxLineLength=110))
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                      bottomMargin=2 * cm, title=pdf_safe(report.title)).build(story)
    return buffer.getvalue()
